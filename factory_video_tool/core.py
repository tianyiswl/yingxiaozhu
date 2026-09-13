"""Single-video domain and cancellable process boundary."""
import json
import math
import random
import re
import shutil
import subprocess
import sys
import threading
import time
from pathlib import Path

class Cancelled(Exception):
    pass

class Runner:
    def __init__(self, cancel=None, log=None):
        self.cancel=cancel or threading.Event()
        self.log=log or (lambda s: None)
    def check(self):
        if self.cancel.is_set(): raise Cancelled('已取消；输入和过程文件已保留')
    def run(self,args,cwd=None,timeout=600):
        self.check()
        args=[str(a) for a in args]
        self.log(json.dumps(args,ensure_ascii=False))
        p=subprocess.Popen(args,cwd=cwd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,
                           creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        start=time.monotonic()
        try:
            while True:
                self.check()
                if time.monotonic()-start>timeout: raise RuntimeError('操作超时：'+Path(args[0]).name)
                try:
                    out,err=p.communicate(timeout=.1);break
                except subprocess.TimeoutExpired: pass
            self.check()
            if p.returncode:
                detail=err.decode('utf-8',errors='replace')
                self.log(detail)
                raise RuntimeError(Path(args[0]).name+' 执行失败：'+detail[-1800:])
            return out.decode('utf-8',errors='replace')
        finally:
            if p.poll() is None:
                p.terminate()
                try:p.communicate(timeout=2)
                except subprocess.TimeoutExpired:p.kill();p.communicate()


def executable(name):
    import os
    bundled_root=getattr(sys,'_MEIPASS',None)
    bundled=Path(bundled_root)/'tools'/(name+('.exe' if os.name=='nt' else '')) if bundled_root else None
    value=os.environ.get('P009_'+name.upper()) or (str(bundled) if bundled and bundled.is_file() else None) or shutil.which(name)
    if not value:raise RuntimeError('缺少 '+name+'；请配置 P009_'+name.upper()+' 的可执行文件路径')
    return value

def probe(path,runner):
    p=Path(path)
    if not p.is_file():raise ValueError('文件不存在：'+str(p))
    return json.loads(runner.run([executable('ffprobe'),'-v','error','-show_format','-show_streams','-of','json',p],timeout=60))

def duration(info):
    d=float(info.get('format',{}).get('duration',0))
    if not math.isfinite(d) or d<=0:raise ValueError('无法识别有效时长')
    return d

def validate_rows(rows):
    result=[];seen=set()
    for index,row in enumerate(rows,2):
        if not any(v is not None and str(v).strip() for v in row.values()):continue
        item={k:str(row.get(k) if row.get(k) is not None else '').strip() for k in ('voice_text','title','body','tags')}
        if not item['voice_text'] or not item['title']:raise ValueError('第 %s 行：标题和口播文案必填'%index)
        enabled=row.get('enabled')
        if enabled in (None,'',1,'1',True):enabled=True
        elif enabled in (0,'0',False):enabled=False
        else:raise ValueError('第 %s 行：是否启用只能填 0 或 1'%index)
        ident=str(row.get('external_id') or row.get('content_id') or 'row-'+str(index)).strip()
        if ident in seen:raise ValueError('内容 ID 重复：'+ident+'；请在 Excel 中明确修改后重新导入')
        seen.add(ident);item.update(external_id=ident,enabled=enabled);result.append(item)
    if not result:raise ValueError('内容包为空')
    return result

EXCEL_HEADER_ALIASES={
    '文案编号':'external_id','内容编号':'content_id','标题':'title','口播文案':'voice_text',
    '发布正文':'body','标签':'tags','是否启用':'enabled',
    'external_id':'external_id','content_id':'content_id','title':'title','voice_text':'voice_text',
    'body':'body','tags':'tags','enabled':'enabled',
}

def load_excel(path):
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        rows=book.active.iter_rows(values_only=True)
        header=next(rows,())
        names=[str(h).strip() if h is not None else '' for h in header]
        names=[EXCEL_HEADER_ALIASES.get(name,name) for name in names]
        if not {'voice_text','title'}.issubset(names):raise ValueError('Excel 首行必须包含“标题”和“口播文案”')
        if len([n for n in names if n])!=len(set(n for n in names if n)):raise ValueError('Excel 列名重复')
        return validate_rows([dict(zip(names,r)) for r in rows])
    finally:book.close()

def subtitle_cues(text,total,max_chars=14):
    if type(max_chars)!=int or not 6<=max_chars<=30:raise ValueError('每句字幕最多显示6到30个字')
    parts=[]
    for sentence in re.split(r'[，。！？；,!?;\n]+',text):
        sentence=sentence.strip()
        parts.extend(sentence[i:i+max_chars] for i in range(0,len(sentence),max_chars))
    if not parts:raise ValueError('口播必须包含可显示的文字')
    weight=sum(map(len,parts));elapsed=0;cues=[]
    for p in parts:
        end=elapsed+total*len(p)/weight;cues.append((elapsed,end,p));elapsed=end
    cues[-1]=(cues[-1][0],total,cues[-1][2]);return cues

def scan_materials(root,runner):
    root=Path(root).resolve()
    if not root.is_dir():raise ValueError('产品素材目录不存在：'+str(root))
    valid=[];errors=[]
    for p in sorted(root.rglob('*')):
        runner.check()
        if p.suffix.lower() not in {'.mp4','.mov','.mkv','.avi'}:continue
        if root not in p.resolve().parents:errors.append('跳过目录外符号链接：'+str(p));continue
        try:
            info=probe(p,runner)
            if not any(s['codec_type']=='video' for s in info['streams']):raise ValueError('没有视频流')
            # Probe alone does not detect damaged frames.
            runner.run([executable('ffmpeg'),'-v','error','-xerror','-i',p,'-map','0:v:0','-an','-f','null','-'],timeout=600)
            valid.append((p,duration(info)))
        except Cancelled:raise
        except Exception as e:errors.append(str(e))
    if not valid:raise ValueError('没有可用视频素材。'+'；'.join(errors)[-1500:])
    return valid,errors

def choose_clips(materials,target):
    if not materials:raise ValueError('素材不足：没有可用视频')
    if target<=0 or any(d<=0 for _,d in materials):raise ValueError('素材或目标时长无效')
    pool=list(materials);random.shuffle(pool);clips=[];remaining=target;i=0
    while remaining>1e-7:
        p,d=pool[i%len(pool)];length=min(random.uniform(3,7),d,remaining)
        start=random.uniform(0,max(0,d-length));clips.append((p,start,length));remaining-=length;i+=1
    return clips

def synthesize_with_retry(provider,request):
    for attempt in range(2):
        try:return provider.synthesize(request)
        except Cancelled:raise
        except Exception:
            if attempt:raise
