"""Create the editable Excel template used by the content library."""
from pathlib import Path

HEADERS=('文案编号','标题','口播文案','发布正文','标签','是否启用')
EXAMPLE_ROW=(
    'tutorial-001',
    '映小助：一分钟上手教程',
    '一分钟，学会映小助。进入账号与设置，点击登录。首次按提示扫码，已有登录会自动核对。打开素材库，点击添加文件夹。粘贴视频文件夹路径，点击导入，整批素材就进来了。再进文案库，点击批量添加文档。先按模板填写标题、口播和标签，再导入文档。回到每日工作台，设置每天制作数量和天数。选择定时发布，设好日期、时间和发布间隔，再点一键制作并安排发布。软件会自动混剪、配音并提交。任务显示平台已接收，就等抖音按定时发布。',
    '一分钟了解映小助的主要操作：登录抖音、文件夹导入素材、导入文案、一键制作并安排定时发布。',
    '软件教程 映小助 使用教程',
    1,
)


def write_content_template(destination):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path=Path(destination).expanduser().resolve()
    if path.suffix.lower()!='.xlsx':path=path.with_suffix('.xlsx')
    path.parent.mkdir(parents=True,exist_ok=True)
    book=Workbook();sheet=book.active;sheet.title='文案'
    sheet.append(HEADERS);sheet.append(EXAMPLE_ROW);sheet.freeze_panes='A2';sheet.auto_filter.ref='A1:F1'
    fill=PatternFill('solid',fgColor='FCE4D6')
    notes={
        '文案编号':'每条唯一编号，例如 factory-001。留空也可自动生成。',
        '标题':'视频标题，必填，最多 30 个字。',
        '口播文案':'配音和字幕内容，必填。可直接粘贴完整口播。',
        '发布正文':'抖音发布正文。留空时可填写与标题相同的简短说明。',
        '标签':'话题用空格或逗号隔开，例如 工厂日常 产品展示。不要写 #。',
        '是否启用':'填 1 导入，填 0 暂不导入。默认填 1。',
    }
    widths=(18,28,56,42,28,12)
    for cell,width in zip(sheet[1],widths):
        cell.font=Font(bold=True,color='9E480E');cell.fill=fill;cell.alignment=Alignment(horizontal='center')
        cell.comment=Comment(notes[cell.value],'映小助')
        sheet.column_dimensions[cell.column_letter].width=width
    for cell in sheet[2]:
        cell.alignment=Alignment(wrap_text=True,vertical='top')
        cell.fill=PatternFill('solid',fgColor='FFF7ED')
    sheet.row_dimensions[2].height=220
    validation=DataValidation(type='list',formula1='"1,0"',allow_blank=False)
    sheet.add_data_validation(validation);validation.add('F2:F1000')
    guide=book.create_sheet('填写说明')
    guide.append(['映小助文案导入模板'])
    guide.append(['第 2 行是程序使用教程的口播示例。可以直接改成自己的内容，或删除后从第 3 行开始填写。'])
    guide.append(['填写完成后，只保留“文案”工作表，按原格式保存，再在文案库中批量添加该文件。'])
    guide.append(['标题与口播文案必填；标签不要输入 #；一行就是一条视频文案。'])
    guide.column_dimensions['A'].width=100
    for row in guide.iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    book.save(path);book.close()
    return path
