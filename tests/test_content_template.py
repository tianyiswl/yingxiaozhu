import tempfile
import unittest
from pathlib import Path

from factory_video_tool.content_template import EXAMPLE_ROW,HEADERS,write_content_template
from factory_video_tool.core import load_excel


class ContentTemplateTests(unittest.TestCase):
    def test_excel_template_has_import_headers_and_a_separate_guide(self):
        with tempfile.TemporaryDirectory() as directory:
            path=write_content_template(Path(directory)/'文案模板')
            from openpyxl import load_workbook
            book=load_workbook(path,data_only=True)
            try:
                self.assertEqual(path.name,'文案模板.xlsx')
                self.assertEqual(book.sheetnames,['文案','填写说明'])
                self.assertEqual(tuple(cell.value for cell in book['文案'][1]),HEADERS)
                self.assertEqual(tuple(cell.value for cell in book['文案'][2]),EXAMPLE_ROW)
                self.assertIn('程序使用教程的口播示例',book['填写说明']['A2'].value)
                self.assertIn('一行就是一条视频文案',book['填写说明']['A4'].value)
            finally:book.close()
            content=load_excel(path)
            self.assertEqual(content[0]['title'],EXAMPLE_ROW[1])
            self.assertEqual(content[0]['voice_text'],EXAMPLE_ROW[2])
            self.assertTrue(content[0]['enabled'])
